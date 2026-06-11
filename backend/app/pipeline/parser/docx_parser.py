from docx import Document
from pathlib import Path
from backend.app.utils.logger import get_logger

logger = get_logger(__name__)


def extract_text_from_docx(file_path: str) -> dict:
    """
    Extracts structured text from a Word (.docx) file.
    Captures paragraphs, headings, and table content.
    Returns a dict with raw text and structured sections.
    """
    result = {
        "raw_text": "",
        "sections": [],
        "tables": [],
    }

    try:
        doc = Document(file_path)
        full_text_parts = []
        current_section = {"heading": "Introduction", "content": []}

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name.lower()

            # Detect headings — start a new section
            if "heading" in style_name:
                # Save previous section if it has content
                if current_section["content"]:
                    result["sections"].append(current_section)
                current_section = {"heading": text, "content": []}
                full_text_parts.append(f"\n## {text}\n")
            else:
                current_section["content"].append(text)
                full_text_parts.append(text)

        # Append the last section
        if current_section["content"]:
            result["sections"].append(current_section)

        # Extract table content
        for table_idx, table in enumerate(doc.tables):
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                row_data = [cell for cell in row_data if cell]
                if row_data:
                    table_data.append(row_data)
                    full_text_parts.append(" | ".join(row_data))

            if table_data:
                result["tables"].append({
                    "table_index": table_idx,
                    "rows": table_data,
                })

        result["raw_text"] = "\n".join(full_text_parts)

        logger.info(
            f"DOCX parsed: {Path(file_path).name} | "
            f"{len(result['sections'])} sections | "
            f"{len(result['tables'])} tables | "
            f"{len(result['raw_text'])} chars"
        )

    except Exception as e:
        logger.error(f"Failed to parse DOCX {file_path}: {e}")
        raise

    return result